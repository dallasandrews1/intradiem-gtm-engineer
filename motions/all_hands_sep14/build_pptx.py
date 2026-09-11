#!/usr/bin/env python3
"""Fill marketing's AI-Champions-Demo template for the GTM Engineering segment.

Source template: ~/Downloads/AI-Champions-Demo (1).pptx (Chelsea / Rachel, Sep 10 2026).
Video: ~/Desktop/Intradiem_AllHands_Demo_v3_Sep9.mp4 (silent by design, narrated live).
Output: ~/Desktop/Intradiem Deliverables/All-hands demo (Sep 10)/AI-Champions-Demo_GTM-Engineering.pptx

Slide order in the output:
  1  AI Champion intro, two presenters (Naveen sets the stage, Dallas runs the demo)
  2  What problem you were trying to solve with AI
  3  Demo video, full slide
  4  Results
  5  Questions / Thank you
"""
import copy, json, os, re, subprocess, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

HOME = Path.home()
SRC = HOME / "Downloads" / "AI-Champions-Demo (1).pptx"
VIDEO = HOME / "Desktop" / "Intradiem_AllHands_Demo_v3_Sep9.mp4"
OUT_DIR = HOME / "Desktop" / "Intradiem Deliverables" / "All-hands demo (Sep 10)"
OUT = OUT_DIR / "AI-Champions-Demo_GTM-Engineering.pptx"
SCRATCH = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/tmp/pptx_build")
SCRATCH.mkdir(parents=True, exist_ok=True)
MEDIA = SCRATCH / "x" / "ppt" / "media"          # unzipped template media (icon library)
HERE = Path(__file__).resolve().parent
CUES = HERE / "final_cut" / "cues.json"

GREEN = RGBColor(0x2C, 0xB5, 0x6E)   # theme accent1
FOREST = RGBColor(0x01, 0x46, 0x37)  # theme accent5
ORANGE = RGBColor(0xF5, 0x82, 0x20)  # theme accent2
INK = RGBColor(0x36, 0x36, 0x36)     # theme dk1
CARD = RGBColor(0xF2, 0xF2, 0xF2)    # bg1 at 95%
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SHADOW_XML = (
    '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:outerShdw blurRad="50800" dist="38100" dir="2700000" algn="tl" rotWithShape="0">'
    '<a:prstClr val="black"><a:alpha val="20000"/></a:prstClr></a:outerShdw></a:effectLst>'
)

# ----------------------------------------------------------------------------- helpers

def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)


def set_text(shape, text, size=None, bold=None, color=None):
    tf = shape.text_frame
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            if size: r.font.size = Pt(size)
            if bold is not None: r.font.bold = bold
            if color: r.font.color.rgb = color
    return tf


def add_paras(tf, items, size=16, gap=6):
    """items: list of (lead, rest). First paragraph reuses the existing one."""
    first = True
    for lead, rest in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if first:
            for r in list(p.runs): r._r.getparent().remove(r._r)
        first = False
        p.level = 0
        p.space_after = Pt(gap)
        if lead:
            r = p.add_run(); r.text = lead; r.font.bold = True; r.font.size = Pt(size)
        if rest:
            r = p.add_run(); r.text = rest; r.font.size = Pt(size)


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def card(slide, x, y, w, h, fill=CARD):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = 0.10
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    spPr = s._element.spPr
    spPr.append(etree.fromstring(SHADOW_XML))
    s.text_frame.text = ""
    return s


def textbox(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing: p.line_spacing = line_spacing
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb


def icon(slide, name, x, y, size=Inches(0.42)):
    return slide.shapes.add_picture(str(MEDIA / name), x, y, height=size)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ----------------------------------------------------------------------------- live counts

def count_maps():
    """Back-office account maps built: Inger's 12 (per-account pages in account_maps/, all built in
    Sales Nav as of Aug 27) plus every other rep set's accounts list (Nate 6, JW 4, Centene/Rachel 1)."""
    sets = HERE.parent / "back_office_expansion" / "sets"
    n = 12
    for f in sorted(sets.glob("*.json")):
        if f.stem in ("inger", "nate_front"):   # nate_front is the front-office view of Nate's same six
            continue
        d = json.load(open(f))
        n += len(d.get("accounts") or [])
    return n


# ----------------------------------------------------------------------------- build

prs = Presentation(str(SRC))
SW, SH = prs.slide_width, prs.slide_height
s_single, s_two, s_problem, s_video, s_results, s_thanks = list(prs.slides)[:6]

TITLE = "GTM Engineering: using AI to widen the net"

# --- Slide: two presenters (becomes slide 1)
s_two.shapes.title.text = TITLE
set_text(ph(s_two, 12), "Naveen Thilagan")
set_text(ph(s_two, 16), "Director, Product Strategy")
set_text(ph(s_two, 17), "Product Strategy")
set_text(ph(s_two, 18), "Leads product strategy. GTM engineering sits on his team.")
set_text(ph(s_two, 19), "Dallas Andrews")
set_text(ph(s_two, 21), "GTM Engineer")
set_text(ph(s_two, 22), "Product Strategy, GTM Engineering")
set_text(ph(s_two, 23), "Joined in July. Builds the AI that does the sourcing, research, and first drafts behind our outbound, so reps start with a warm list instead of a blank page.")

# --- Slide: single presenter. Both Naveen and Dallas present (DM Sep 10), so this slide is removed.
sldIdLst = prs.slides._sldIdLst
single_id = [e for e in sldIdLst if prs.slides.get(int(e.get("id"))) is s_single][0]
prs.part.drop_rel(single_id.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
sldIdLst.remove(single_id)

# --- Slide: problem (marketing's pre-written question stays as the title)
s_problem.shapes.title.text = "The problem we're solving with AI"
set_text(ph(s_problem, 15),
         "Three versions of one problem: the work that has to happen before a seller can start a conversation.",
         size=18)
left = ph(s_problem, 12).text_frame
add_paras(left, [
    ("The problem", ""),
    ("A target account and a blank page. ", "A day of research per account before the first message goes out."),
    ("Our customers' back offices don't know us. ", "Claims, billing, enrollment, payments. Nobody in the front office introduces us."),
    ("Five touches across three channels for every contact. ", "A full-time job before a single reply comes back."),
    ("About 1,800 hours of manual work. ", "By hand, it just doesn't happen."),
], size=18)
right = ph(s_problem, 14).text_frame
add_paras(right, [
    ("What we set out to do", ""),
    ("Find the people worth reaching, ", "outside Salesforce and inside our own customers' back offices."),
    ("Research every contact and write the first draft ", "in the rep's voice, using only numbers we can back up."),
    ("Run the sequences automatically, ", "with a person deciding at the points that matter. Reps own every send and every reply."),
    ("Count it in Salesforce ", "like any other channel: cost per meeting, on the same dashboard leadership already reads."),
], size=18)
# make the two column headers green
for tf in (left, right):
    p0 = tf.paragraphs[0]
    r = p0.runs[0]
    r.font.color.rgb = GREEN; r.font.size = Pt(20)
    pPr = p0._p.get_or_add_pPr()
    pPr.set("indent", "0"); pPr.set("marL", "0")
    pPr.insert(0, etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buNone"))

# --- Slide: demo video, full slide
for shp in list(s_video.shapes):
    remove_shape(shp)
poster = SCRATCH / "poster.png"
if not poster.exists():
    subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-ss", "1.5", "-i", str(VIDEO),
                    "-frames:v", "1", str(poster)], check=True)
movie = s_video.shapes.add_movie(str(VIDEO), 0, 0, SW, SH, poster_frame_image=str(poster), mime_type="video/mp4")
movie.name = "GTM Engineering demo (4:57)"
for rid, rel in list(s_video.part.rels.items()):      # drop the template's dangling recording-guide link
    if rel.is_external:
        s_video.part.drop_rel(rid)

# --- Slide: results
for shp in list(s_results.shapes):
    if shp != s_results.shapes.title:
        remove_shape(shp)
s_results.shapes.title.text = "Results, ten weeks in"

maps_n = count_maps() or 18
x0, gap = Inches(0.5), Inches(0.2)
cw = (SW - 2 * x0 - 3 * gap) / 4
# row 1: the numbers
r1y, r1h = Inches(1.42), Inches(2.0)
row1 = [
    ("image42.png", "370 hours", "of sourcing and research handed back to sellers"),
    ("image35.png", "4,600", "new contacts found outside Salesforce"),
    ("image55.png", "12.4K", "existing contacts segmented at no cost"),
    ("image54.png", "950", "accounts scored, checked, or mapped"),
]
for i, (ic, num, desc) in enumerate(row1):
    x = x0 + i * (cw + gap)
    c = card(s_results, x, r1y, cw, r1h, fill=(FOREST if i == 0 else CARD))
    icon(s_results, ic, x + cw - Inches(0.62), r1y + Inches(0.18))
    textbox(s_results, x + Inches(0.15), r1y + Inches(0.55), cw - Inches(0.3), Inches(0.75), num, 34, bold=True,
            color=(WHITE if i == 0 else GREEN))
    textbox(s_results, x + Inches(0.15), r1y + Inches(1.3), cw - Inches(0.3), Inches(0.65), desc, 14,
            color=(WHITE if i == 0 else INK))
# row 2: what is live
r2y, r2h = Inches(3.62), Inches(1.62)
row2 = [
    ("image36.png", "8 campaigns sending", "Live in North America and the UK. Every reply goes straight to the rep."),
    ("image31.png", f"{maps_n} account maps", "Back-office org charts for the AM team, every name checked against Salesforce."),
    ("image44.png", "24 agents on a schedule", "Overnight scans, audits, and briefs waiting in Slack each morning."),
    ("image30.png", "ICP buying committees", "Four roles per account instead of one contact. Now how every campaign targets."),
]
for i, (ic, label, desc) in enumerate(row2):
    x = x0 + i * (cw + gap)
    card(s_results, x, r2y, cw, r2h)
    icon(s_results, ic, x + Inches(0.15), r2y + Inches(0.15), size=Inches(0.4))
    textbox(s_results, x + Inches(0.62), r2y + Inches(0.1), cw - Inches(0.75), Inches(0.5), label, 16, bold=True,
            color=FOREST, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s_results, x + Inches(0.15), r2y + Inches(0.65), cw - Inches(0.3), Inches(0.8), desc, 13)
# row 3: the time-saved line
r3y, r3h = Inches(5.44), Inches(1.15)
c = card(s_results, x0, r3y, SW - 2 * x0, r3h, fill=WHITE)
bar = s_results.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0, r3y + Inches(0.15), Inches(0.08), r3h - Inches(0.3))
bar.fill.solid(); bar.fill.fore_color.rgb = ORANGE; bar.line.fill.background(); bar.shadow.inherit = False
textbox(s_results, x0 + Inches(0.3), r3y + Inches(0.1), SW - 2 * x0 - Inches(0.5), Inches(0.45),
        "Time saved", 15, bold=True, color=ORANGE)
textbox(s_results, x0 + Inches(0.3), r3y + Inches(0.5), SW - 2 * x0 - Inches(0.5), Inches(0.6),
        "An account plan that used to take a rep a day now takes one ask and a few minutes. An org map that took an AM a week now takes days. "
        "That time goes back into conversations, and every meeting gets counted in Salesforce like any other channel.",
        13.5)

# --- Slide: questions
set_text(ph(s_thanks, 12), "Questions?")
set_text(ph(s_thanks, 13), "Thank you!")

# ----------------------------------------------------------------------------- speaker notes
cues = json.load(open(CUES))

def mmss(t):
    return f"{int(t)//60}:{int(t)%60:02d}"

narration = "\n".join(f"{mmss(c['t'])}  {c['line']}" for c in cues)

notes(s_two,
      "NAVEEN (about 2 min): why the company is hearing this. Go-to-market is being orchestrated through AI: the same engine "
      "finds the people, does the research, writes the first drafts, and runs the sequences, with people deciding at the "
      "points that matter. The one idea to watch for: we used to sell to a person, now we sell to the buying committee. "
      "Then hand to Dallas.\n\n"
      "DALLAS (30 sec): name, role, one line. 'Same problem three times, solved end to end. Here is what it looks like.'")
notes(s_problem,
      "DALLAS (about 60 sec). Read the three problems as the room already knows them, then the right column as the bar we set.\n\n"
      "Frame it as time given back, never people replaced: by hand this work never happens, so the hours the engine does go back to sellers as conversation time.\n\n"
      "If pressed on 1,800 hours: an estimate. 4,400 contacts (the count when the estimate was made) at roughly 25 minutes each "
      "to find, verify, research, and write five touches by hand. Say it is an estimate.")
notes(s_video,
      "DEMO, 4:57. The video is silent by design; Dallas narrates live from the lines below (Naveen, Sep 9: 'Please do it live'). "
      "Click the frame to play. Timecodes are minutes:seconds into the video.\n\n" + narration)
notes(s_results,
      "DALLAS (about 90 sec). Lead with time saved, then the numbers, then what is live today.\n\n"
      "If pressed on 370 hours: 4,400 contacts at 5 minutes each of sourcing and research alone, the most conservative cut. An estimate, say so.\n"
      "Counts as of Sep 10 2026: 4,600 contacts sourced outside Salesforce; 12.4K existing Salesforce contacts segmented in Clay at zero cost; "
      "950 accounts scored, gated, or mapped; 8 campaigns running in lemlist (Nate 4 NA, Jack 4 UK); "
      f"{maps_n} back-office account maps across the AM and sales rep sets; 24 scheduled agents (launchd).\n\n"
      "Frame: time given back to sellers, never headcount. Keep estimated pipeline and real results separate; real only counts once it is in Salesforce.")
notes(s_thanks,
      "Likely questions, one line each:\n"
      "Tools: Claude Code is the brain, Clay the data layer, lemlist runs sends, Salesforce stays the system of record.\n"
      "Does it send by itself: it writes, sequences, and runs at scale, with a person deciding at the points that matter, tighter on new campaigns; current customers never enter a cold campaign; only numbers we can back up go into outreach.\n"
      "Accuracy: every claim has to exist in the verified value repository with a source, or it doesn't go in.\n"
      "Replacing sellers: no. It removes sourcing, research, and first-draft work; reps own every send, reply, and meeting.\n"
      "Could my team use it: yes, it already runs for account management and alliances; it is built as self-serve skills teams run themselves.\n"
      "Security: data stays inside tools the company already controls; new AI tools go through the contracts-team review.\n"
      "Anything else: answer what is true today in one sentence, then 'grab me after, I'll show you the real thing'.")

OUT_DIR.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print("saved", OUT, f"{OUT.stat().st_size/1e6:.1f} MB", "maps", maps_n)
