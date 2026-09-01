#!/usr/bin/env python3
"""Intake for the 3xG workbook "Intradiem Work -Verint Partner Account Whitespace 2024.xlsx"
(Jeremy Roderick, 3xG Consulting, forwarded by Frank Ciccone, Aug 31 2026).

Fitted to the workbook as received (two sheets):

  "Verint Partner Only 2024"          74 partner-sold Verint accounts:
      Partner (the Verint reseller of record, e.g. Avaya, ConvergeOne, Cisco, AWS), Account,
      Max Verint (size band), Max Verint license (seats), Model, DPA Seats, Vertical Alignment
      (1/2/3, scale as given), DPA-Max Licence (seats without desktop analytics, the DPA whitespace)
    plus a bottom block of 9 Direct/TTEC rows: Account, back-office product (Operations Visualizer,
      Operations Manager, Operations Productivity, eg Work Manager), seats.

  "Quantified Whitespac No SSA-IRS"   1,297 Verint accounts (the whole whitespace matrix):
      Bill To Name, Account Name, Bucket (size band), Max License, Type (WFE), per-module license
      counts (Rec, QM, DPA, WFM, WFMPro, PM, Speech ...), then dollar whitespace per add-on and a Total.

What it does, read-only on every source, 0 credits:
  1. Buckets each account the way Frank described the list on the call:
        operational_visualizer   Verint Operations Visualizer / Operations Manager / Operations
                                 Productivity deployed (a first back-office deployment, then nothing):
                                 Frank's priority
        backoffice_deployment    Verint eg Work Manager deployed
        dpa_deployed             desktop analytics in place (DPA Seats > 0)
        verint_no_dpa            Verint front-office footprint, no desktop analytics
  2. Pulls two facts from the big sheet for each partner-sold account: whether WFM or WFM Pro is in
     the footprint (Intradiem sits on WFM) and the Verint seat count when the partner sheet has none.
  3. STRIPS the 3xG/Verint-sensitive columns: every dollar whitespace column, list prices, totals,
     and the per-module license counts. Only the fields listed in OUTPUT_COLUMNS leave the intake.
     A scrub manifest records exactly what was dropped.
  4. Runs the customer-exclusion gate against every source on disk (Active_Customers_SF_Jul10.csv,
     customer_denylist.json aliases and domains, UK current customers, and the Salesforce
     Customer-or-Partner segment snapshot when one has been pulled into ./ref/). Matches are HELD for
     the account manager and never ranked into a cold cohort. Partners (Five9 etc.) are held too.
  5. Ranks: bucket, then seats (log scale), vertical alignment as given, WFM footprint, TAM match.
  6. Writes, next to the workbook in run_<date>/:
        3xg_candidates_ranked.csv    the pilot universe (partner-sold + direct back-office rows)
        3xg_extended_pool.csv        big-sheet accounts, 500+ seats, WFM or DPA present, not on the
                                     partner sheet (for Frank if the cohort needs more choice)
        3xg_intake_report.md         what was read, dropped, held, and the top 15, plus questions for Frank
        3xg_scrub_manifest.json

Run:
  python3 intake_3xg_whitespace.py                  # newest .xlsx in the inbox
  python3 intake_3xg_whitespace.py --file path.xlsx
  python3 intake_3xg_whitespace.py --selftest       # synthetic two-sheet workbook, proves the pipe
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import math
import os
import pathlib
import re
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: python3 -m pip install openpyxl")

HERE = pathlib.Path(__file__).resolve().parent
REPO = HERE.parent.parent
MAIN_REPO = pathlib.Path.home() / "Claude" / "Projects" / "Intradiem GTM Engineer"
DEFAULT_INBOX = MAIN_REPO / "automation" / "inbox" / "partners" / "3xg"

TAM_CSV = REPO / "tam-outbound-engine" / "data" / "tam_accounts.csv"
DENYLIST = REPO / "tam-outbound-engine" / "config" / "customer_denylist.json"
SF_CUSTOMERS = REPO / "greenlight-pack" / "Active_Customers_SF_Jul10.csv"
UK_CUSTOMERS = REPO / "motions" / "UK_Current_Customers_Aug3.csv"
REF_DIR = HERE / "ref"  # optional snapshots, e.g. sf_customer_partner_segment_<date>.csv (name,domain,type)

PARTNER_SHEET = "Verint Partner Only 2024"
MATRIX_SHEET = "Quantified Whitespac"

OUTPUT_COLUMNS = [
    "rank", "account", "source_list", "reseller_of_record", "bucket", "size_band", "verint_seats",
    "dpa_seats", "dpa_whitespace_seats", "model", "back_office_product", "vertical_alignment_as_given",
    "wfm_in_footprint", "flags", "score",
]
BUCKET_SCORE = {"operational_visualizer": 100, "backoffice_deployment": 80, "dpa_deployed": 60, "verint_no_dpa": 30}
VERT_BONUS = {"1": 15, "2": 8, "3": 0}

OV_PAT = re.compile(r"operations?\s*(visuali[sz]er|manager|productivity)|process\s*operations|strategic\s*operations", re.I)
WM_PAT = re.compile(r"work\s*manager", re.I)

# Matrix-sheet columns that never leave the intake (dollar whitespace, prices, totals) or are
# competitive detail we do not export (per-module license counts). WFM/WFMPro/DPA are read for a
# yes/no and a seat count only.
DOLLAR_COLS = ("Workflex", "Quality Bot", "Trans Bot", "Wrap Up", "Voice Capture", "RT Coach", "Redaction", "Speech 480", "Encryption", "Total")
MODULE_COLS = ("EDM-Rec", "Rec", "Encrp", "QM", "AQM", "PM", "Speech", "RTAA", "IQ", "CF", "XM", "VFC", "community", "RecPubSafe", "ChAuto", "KMEnt", "KMPro")

STOP = {"inc", "llc", "ltd", "corp", "corporation", "co", "company", "plc", "group", "holdings", "the", "of", "and", "na", "n", "a", "dba", "fka", "llp", "lp", "sa", "ag", "se"}
# First tokens too generic to identify a brand on their own (a shared first token here is a CHECK, not a HOLD).
GENERIC = {"southern", "northern", "western", "eastern", "central", "american", "america", "national", "first", "united", "general",
           "pacific", "new", "global", "international", "health", "healthcare", "medical", "hospital", "university", "bank", "insurance",
           "financial", "services", "service", "energy", "state", "department", "credit", "union", "mutual", "life", "care", "center",
           "systems", "solutions", "technology", "capital", "trust", "federal", "city", "county", "north", "south", "east", "west", "us", "usa"}
# Brand families that trade under several first tokens.
TOKEN_CANON = {"citibank": "citi", "citicorp": "citi", "citigroup": "citi", "jpmorgan": "chase", "aflac": "aflac", "americanfamily": "aflac"}


def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def key(s) -> str:
    s = norm(s).lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    toks = [t for t in s.split() if t not in STOP]
    return " ".join(toks)


def name_parts(s: str) -> list[str]:
    """Split 'X dba Y', 'X (fka Y)', 'X - Y' into candidate names, keys included."""
    s = norm(s)
    parts = re.split(r"\s+(?:dba|d/b/a|fka|f/k/a|aka)\s+|\(|\)|\s+-\s+|/", s, flags=re.I)
    out = []
    for p in parts:
        k = key(p)
        if k:
            out.append(k)
    return out or [key(s)]


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    m = re.search(r"-?\d[\d,]*\.?\d*", str(v))
    if not m:
        return None
    try:
        return float(m.group(0).replace(",", ""))
    except ValueError:
        return None


def fmt_int(v):
    return "" if v is None else (int(v) if float(v).is_integer() else v)


# ---- reference lists ------------------------------------------------------
class Gate:
    """Customer and partner exclusion, built from every source on disk. Matches hold, never drop."""

    def __init__(self):
        self.entries: dict[str, str] = {}   # key -> source label
        self.names: dict[str, str] = {}     # key -> name as it appears in the source
        self.aliases: list[tuple[str, str]] = []
        self.domains: dict[str, str] = {}
        self.tam: set[str] = set()
        self.sources: list[str] = []
        self._load()

    def _add(self, name, label):
        for k in name_parts(name):
            if len(k) >= 3:
                self.entries.setdefault(k, label)
                self.names.setdefault(k, norm(name))

    def _load(self):
        # Segment snapshot first: it is the most recent Salesforce truth (Customer or Partner), so its labels win.
        if REF_DIR.exists():
            for p in sorted(REF_DIR.glob("sf_customer_partner_segment_*.csv")):
                with p.open(newline="", encoding="utf-8-sig") as f:
                    n = 0
                    for r in csv.DictReader(f):
                        nm = r.get("name") or r.get("company") or r.get("account_name") or ""
                        typ = r.get("type") or "Customer or Partner"
                        if nm:
                            self._add(nm, f"SF {typ} (Clay segment {p.stem[-10:]})")
                            n += 1
                        dom = (r.get("domain") or "").lower()
                        if dom:
                            self.domains.setdefault(dom, f"SF {typ} domain")
                self.sources.append(f"{p.name} ({n})")
        if SF_CUSTOMERS.exists():
            with SF_CUSTOMERS.open(newline="", encoding="utf-8-sig") as f:
                n = 0
                for r in csv.DictReader(f):
                    self._add(r.get("account_name", ""), f"SF {r.get('type') or 'Customer'} (Jul 10 export)")
                    n += 1
            self.sources.append(f"Active_Customers_SF_Jul10.csv ({n})")
        if UK_CUSTOMERS.exists():
            with UK_CUSTOMERS.open(newline="", encoding="utf-8-sig") as f:
                rd = csv.DictReader(f)
                col = next((c for c in (rd.fieldnames or []) if re.search(r"account|company|name", c, re.I)), None)
                n = 0
                for r in rd:
                    if col and r.get(col):
                        self._add(r[col], "UK current customer (Aug 3)")
                        n += 1
            self.sources.append(f"UK_Current_Customers_Aug3.csv ({n})")
        if DENYLIST.exists():
            try:
                d = json.loads(DENYLIST.read_text())
            except json.JSONDecodeError:
                d = {}
            for a in d.get("name_aliases", []):
                self.aliases.append((key(a), "denylist alias"))
            for dom in d.get("domains", []):
                self.domains[dom.lower()] = "denylist domain"
            self.sources.append(f"customer_denylist.json ({len(d.get('name_aliases', []))} aliases, {len(d.get('domains', []))} domains)")
        if TAM_CSV.exists():
            with TAM_CSV.open(newline="", encoding="utf-8-sig") as f:
                for r in csv.DictReader(f):
                    for k in name_parts(r.get("company", "")):
                        self.tam.add(k)
            self.sources.append(f"tam_accounts.csv ({len(self.tam)} names)")

    @staticmethod
    def _phrase_in(needle: str, hay: str) -> bool:
        return re.search(r"(?<![a-z0-9])" + re.escape(needle) + r"(?![a-z0-9])", hay) is not None

    @staticmethod
    def _first(k: str) -> str:
        t = k.split()[0] if k else ""
        return TOKEN_CANON.get(t, t)

    def match(self, account: str):
        """Returns (level, label): level is 'HOLD' (same account or brand) or 'CHECK' (near-name, confirm with the AM), else None."""
        parts = name_parts(account)
        for p in parts:
            if p in self.entries:
                return "HOLD", f"{self.entries[p]}, '{self.names[p]}'"
        for p in parts:
            for a, label in self.aliases:
                if a and self._phrase_in(a, p):
                    return "HOLD", label
        # brand token: same distinctive first token (farmers / farmers insurance, pnc bank / pnc financial, citibank / citicorp)
        for p in parts:
            fp = self._first(p)
            if len(fp) >= 3 and fp not in GENERIC:
                for k, label in self.entries.items():
                    if self._first(k) == fp:
                        return "HOLD", f"{label}, brand match on '{self.names[k]}', AM confirms it is the same company"
        # near-name: one key contains the other as a phrase but the first tokens differ (southern company vs southern california edison)
        for p in parts:
            for k, label in self.entries.items():
                if len(k) >= 6 and (self._phrase_in(k, p) or (len(p) >= 8 and self._phrase_in(p, k))):
                    return "CHECK", f"{label}, '{self.names[k]}'"
        return None

    def reseller_note(self, reseller: str):
        m = self.match(reseller) if reseller else None
        return m[1] if m and m[0] == "HOLD" and "Partner" in m[1] else None

    def in_tam(self, account: str) -> bool:
        return any(p in self.tam for p in name_parts(account))


# ---- workbook ---------------------------------------------------------------
def load_sheets(path: pathlib.Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        out[ws.title] = [r for r in rows if any(c not in (None, "") for c in r)]
    return out


def find_sheet(sheets: dict, needle: str):
    for name in sheets:
        if needle.lower() in name.lower():
            return name, sheets[name]
    return None, None


def parse_matrix(rows):
    """Return {account_key: {...}} from the big whitespace sheet plus the header list."""
    hi = next(i for i, r in enumerate(rows) if r and norm(r[0]).lower() == "bill to name")
    headers = [norm(c) for c in rows[hi]]
    col = {re.sub(r"\s+", " ", h).strip(): i for i, h in enumerate(headers) if h}
    def c(name):
        for k, i in col.items():
            if k.lower().startswith(name.lower()):
                return i
        return None
    i_bill, i_acct, i_bucket, i_max, i_type = c("Bill To"), c("Account Name"), c("Bucket"), c("Max License"), c("Type")
    i_dpa, i_wfm, i_wfmpro = c("DPA"), c("WFM"), c("WFMPro")
    # exact WFM column (not WFMPro)
    i_wfm = next((i for k, i in col.items() if k.upper() == "WFM"), i_wfm)
    acc = {}
    for r in rows[hi + 1:]:
        if len(r) <= i_acct or not norm(r[i_acct]):
            continue
        name = norm(r[i_acct])
        if name.lower() in ("account name", "totals"):
            continue
        rec = {
            "account": name, "bill_to": norm(r[i_bill]) if i_bill is not None else "",
            "band": norm(r[i_bucket]) if i_bucket is not None else "",
            "max_license": to_num(r[i_max]) if i_max is not None else None,
            "model": norm(r[i_type]) if i_type is not None else "",
            "dpa": to_num(r[i_dpa]) if i_dpa is not None else None,
            "wfm": (to_num(r[i_wfm]) or 0) + (to_num(r[i_wfmpro]) or 0) if i_wfm is not None else None,
        }
        for k in name_parts(name):
            acc.setdefault(k, rec)
        acc.setdefault(key(name), rec)
    return acc, headers


def parse_partner(rows):
    hdr = [norm(c) for c in rows[0]]
    def idx(pref):
        return next((i for i, h in enumerate(hdr) if h.lower().startswith(pref.lower())), None)
    i_p, i_a, i_band, i_lic, i_model, i_dpa, i_vert, i_ws = (idx("Partner"), idx("Account"), idx("Max Verint"), idx("Max Verint license"), idx("Model"), idx("DPA Seats"), idx("Vertical"), idx("DPA-Max"))
    # "Max Verint" matches both band and license columns; take the first as band and the explicit one as license
    if i_band == i_lic:
        i_band = next((i for i, h in enumerate(hdr) if h.lower() == "max verint"), i_band)
    partner_rows, direct_rows = [], []
    for r in rows[1:]:
        r = list(r) + [None] * (len(hdr) - len(r))
        acct = norm(r[i_a]) if i_a is not None else ""
        if not acct:
            continue
        partner = norm(r[i_p]) if i_p is not None else ""
        model = norm(r[i_model]) if i_model is not None else ""
        band = norm(r[i_band]) if i_band is not None else ""
        lic = to_num(r[i_lic]) if i_lic is not None else None
        if not band and (OV_PAT.search(model) or WM_PAT.search(model) or partner.lower() in ("direct", "ttec")):
            # bottom block: Account, product in the Model column, seats in the DPA Seats column
            seats = to_num(r[i_dpa]) if i_dpa is not None else None
            direct_rows.append({"account": acct, "partner": partner, "product": model, "seats": seats})
        else:
            partner_rows.append({
                "account": acct, "partner": partner, "band": band, "seats": lic, "model": model,
                "dpa": to_num(r[i_dpa]) if i_dpa is not None else None,
                "vert": norm(r[i_vert]) if i_vert is not None else "",
                "dpa_ws": to_num(r[i_ws]) if i_ws is not None else None,
            })
    return partner_rows, direct_rows, hdr


def score_row(bucket, seats, vert, wfm_yes, tam):
    s = BUCKET_SCORE[bucket]
    if seats and seats > 0:
        s += min(30.0, 8 * math.log10(seats))
    s += VERT_BONUS.get(str(vert).strip(), 0)
    s += 10 if wfm_yes else 0
    s += 10 if tam else 0
    return round(s, 1)


def run(path: pathlib.Path, outdir: pathlib.Path) -> dict:
    sheets = load_sheets(path)
    pname, prows = find_sheet(sheets, PARTNER_SHEET)
    mname, mrows = find_sheet(sheets, MATRIX_SHEET)
    if not prows:
        sys.exit(f"sheet containing '{PARTNER_SHEET}' not found; sheets: {list(sheets)}")
    matrix, mheaders = parse_matrix(mrows) if mrows else ({}, [])
    partner_rows, direct_rows, phdr = parse_partner(prows)
    gate = Gate()

    def lookup(acct):
        for k in name_parts(acct):
            if k in matrix:
                return matrix[k]
        return None

    out = []
    for r in partner_rows:
        m = lookup(r["account"])
        wfm_yes = bool(m and m["wfm"])
        seats = r["seats"] or (m["max_license"] if m else None)
        bucket = "dpa_deployed" if (r["dpa"] or 0) > 0 else "verint_no_dpa"
        gm = gate.match(r["account"])
        held = gm[1] if gm and gm[0] == "HOLD" else None
        flags = []
        if held:
            flags.append(f"HOLD: {held}")
        elif gm:
            flags.append(f"CHECK near-name: {gm[1]}, confirm not related before outreach")
        if gate.in_tam(r["account"]):
            flags.append("in TAM list")
        if r["partner"]:
            rn = gate.reseller_note(r["partner"])
            flags.append(f"Verint reseller of record: {r['partner']}" + (" (also an Intradiem partner)" if rn else ""))
        sc = -1 if held else score_row(bucket, seats, r["vert"], wfm_yes, gate.in_tam(r["account"]))
        out.append({
            "rank": "", "account": r["account"], "source_list": "partner_only_2024", "reseller_of_record": r["partner"],
            "bucket": bucket, "size_band": r["band"], "verint_seats": fmt_int(seats), "dpa_seats": fmt_int(r["dpa"]),
            "dpa_whitespace_seats": fmt_int(r["dpa_ws"]), "model": r["model"], "back_office_product": "",
            "vertical_alignment_as_given": r["vert"], "wfm_in_footprint": "yes" if wfm_yes else ("no" if m else "not on matrix sheet"),
            "flags": "; ".join(flags), "score": sc,
        })
    for r in direct_rows:
        m = lookup(r["account"])
        wfm_yes = bool(m and m["wfm"])
        bucket = "operational_visualizer" if OV_PAT.search(r["product"]) else ("backoffice_deployment" if WM_PAT.search(r["product"]) else "verint_no_dpa")
        gm = gate.match(r["account"])
        held = gm[1] if gm and gm[0] == "HOLD" else None
        flags = ["Frank priority" if bucket == "operational_visualizer" else "back-office deployment", f"sold {r['partner']}, not partner-registered"]
        if held:
            flags.insert(0, f"HOLD: {held}")
        elif gm:
            flags.insert(0, f"CHECK near-name: {gm[1]}, confirm not related before outreach")
        if any(p in {k for x in partner_rows for k in name_parts(x["account"])} for p in name_parts(r["account"])):
            flags.append("also on the partner sheet")
        if gate.in_tam(r["account"]):
            flags.append("in TAM list")
        sc = -1 if held else score_row(bucket, r["seats"], "", wfm_yes, gate.in_tam(r["account"]))
        out.append({
            "rank": "", "account": r["account"], "source_list": "direct_backoffice_block", "reseller_of_record": r["partner"],
            "bucket": bucket, "size_band": m["band"] if m else "", "verint_seats": fmt_int(r["seats"] or (m["max_license"] if m else None)),
            "dpa_seats": fmt_int(m["dpa"]) if m else "", "dpa_whitespace_seats": "", "model": m["model"] if m else "",
            "back_office_product": r["product"], "vertical_alignment_as_given": "",
            "wfm_in_footprint": "yes" if wfm_yes else ("no" if m else "not on matrix sheet"),
            "flags": "; ".join(flags), "score": sc,
        })

    out.sort(key=lambda x: (x["score"], x["verint_seats"] if isinstance(x["verint_seats"], (int, float)) else 0), reverse=True)
    n = 0
    for r in out:
        if r["score"] >= 0:
            n += 1
            r["rank"] = n

    # extended pool from the matrix: 500+ seats, WFM or DPA present, not on the partner sheet
    on_partner = {k for r in out for k in name_parts(r["account"])}
    ext, seen = [], set()
    for k, m in matrix.items():
        if id(m) in seen:
            continue
        seen.add(id(m))
        if (m["max_license"] or 0) < 500 or not ((m["wfm"] or 0) > 0 or (m["dpa"] or 0) > 0):
            continue
        if any(p in on_partner for p in name_parts(m["account"])):
            continue
        gm = gate.match(m["account"])
        fl = [f"{gm[0]}: {gm[1]}"] if gm else []
        if gate.in_tam(m["account"]):
            fl.append("in TAM list")
        ext.append({
            "account": m["account"], "reseller_of_record": m["bill_to"], "size_band": m["band"], "verint_seats": fmt_int(m["max_license"]),
            "dpa_seats": fmt_int(m["dpa"]), "wfm_in_footprint": "yes" if m["wfm"] else "no", "model": m["model"],
            "flags": "; ".join(fl),
        })
    ext.sort(key=lambda x: (0 if x["flags"].startswith("HOLD") else 1, x["verint_seats"] if isinstance(x["verint_seats"], (int, float)) else 0), reverse=True)

    outdir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    with (outdir / "3xg_candidates_ranked.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS)
        w.writeheader()
        w.writerows(out)
    with (outdir / "3xg_extended_pool.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(ext[0].keys()) if ext else ["account"])
        w.writeheader()
        w.writerows(ext)

    dropped = [h for h in mheaders if h and (any(h.startswith(d) for d in DOLLAR_COLS) or h in MODULE_COLS)]
    manifest = {
        "source_file": path.name, "run_at": stamp,
        "sheets": {pname: {"rows": len(prows), "headers": phdr}, mname: {"rows": len(mrows or []), "headers": mheaders}},
        "exported_columns": OUTPUT_COLUMNS,
        "dropped_from_matrix_sheet": dropped,
        "dropped_from_partner_sheet": [],
        "read_but_summarised_only": ["WFM", "WFMPro (as wfm_in_footprint yes/no)", "DPA (seat count only)", "Max License (seat count only)"],
        "preamble_rows_ignored": "List Price and Total Opportunity rows above the header on the matrix sheet",
        "gate_sources": gate.sources,
        "rule": "Only exported_columns leave the intake. Dollar whitespace, list prices, totals and per-module license counts never do.",
    }
    (outdir / "3xg_scrub_manifest.json").write_text(json.dumps(manifest, indent=2))

    held = [r for r in out if r["score"] < 0]
    ranked = [r for r in out if r["score"] >= 0]
    counts = {}
    for r in out:
        counts[r["bucket"]] = counts.get(r["bucket"], 0) + 1
    n_matrix = len({v["account"] for v in matrix.values()})
    L = [
        f"# 3xG whitespace intake, {stamp}", "",
        f"Source: `{path.name}`. Sheets read: `{pname}` ({len(partner_rows)} partner-sold rows + {len(direct_rows)} direct back-office rows) and `{mname}` ({n_matrix} distinct accounts, used for WFM footprint and seat fallback). Read-only; no Salesforce, Clay or lemlist call; 0 credits.", "",
        "## Stripped before anything left the intake",
        f"- Matrix sheet: {len(dropped)} columns dropped (every dollar whitespace column, Total, and the per-module license counts): " + ", ".join(f"`{d}`" for d in dropped),
        "- Preamble rows (List Price, Total Opportunity) ignored.",
        f"- Exported fields only: {', '.join(OUTPUT_COLUMNS)}.", "",
        "## Customer and partner gate",
        f"- Sources: {'; '.join(gate.sources)}",
        f"- Held ({len(held)}), never ranked into a cold cohort:",
    ] + ([f"  - {r['account']} ({r['bucket']}, {r['verint_seats'] or 'n/a'} seats): {r['flags'].split(';')[0]}" for r in held] or ["  - none matched"]) + [
        f"- Near-names to confirm with the AM before outreach ({sum(1 for r in ranked if 'CHECK' in r['flags'])}), still ranked:",
    ] + ([f"  - {r['account']}: {r['flags'].split(';')[0]}" for r in ranked if 'CHECK' in r['flags']] or ["  - none"]) + [
        "", "## Buckets (whole universe)",
    ] + [f"- {b}: {counts[b]}" for b in sorted(counts, key=lambda b: -BUCKET_SCORE[b])] + [
        "", "## Top 15 for the Sep 1 review", "",
        "| # | Account | Bucket | Seats | DPA seats | WFM | Vert | Reseller | Back-office product |", "|---|---|---|---|---|---|---|---|---|",
    ] + [f"| {r['rank']} | {r['account']} | {r['bucket']} | {r['verint_seats']} | {r['dpa_seats']} | {r['wfm_in_footprint']} | {r['vertical_alignment_as_given']} | {r['reseller_of_record']} | {r['back_office_product']} |" for r in ranked[:15]] + [
        "", f"Extended pool (matrix accounts with 500+ seats and WFM or DPA present, not on the partner sheet): {len(ext)} rows in `3xg_extended_pool.csv`.", "",
        "## Ranking",
        "Bucket first (operational visualizer 100, back-office deployment 80, desktop analytics deployed 60, Verint footprint without DPA 30), then seats on a log scale (max +30), vertical alignment as given (1: +15, 2: +8), WFM in the footprint (+10), TAM-list match (+10). Held accounts score -1 and carry no rank.", "",
        "## To settle with Frank on Sep 1",
        "- What the Vertical Alignment scale means (1 looks like the strongest fit to our six verticals; confirm before it stays in the ranking).",
        "- The direct back-office block (Operations Visualizer, Work Manager) is Verint-direct, not partner-registered: is it in play for a 3xG motion, and who makes the intro?",
        "- The Partner column is the Verint reseller of record (Avaya, ConvergeOne, Cisco, AWS), not 3xG. Outreach names 3xG only where Jeremy agrees; the reseller is never named.",
        "- Any held account Frank wants worked goes through the account manager, not this pilot.",
    ]
    (outdir / "3xg_intake_report.md").write_text("\n".join(L) + "\n")
    return {"out": str(outdir), "universe": len(out), "ranked": len(ranked), "held": len(held), "buckets": counts, "extended_pool": len(ext), "gate_sources": gate.sources}


def selftest(scratch: pathlib.Path):
    scratch.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quantified Whitespac No SSA-IRS"
    ws.append([None] * 27 + ["Prereq's", "150 min."])
    ws.append([None] * 27 + ["List Price", 240, 240])
    ws.append(["Totals", "Totals"])
    ws.append(["Bill To Name", "Account Name", "Bucket", None, "Max License", "Type    (WFE)", "EDM-Rec", "Rec", "Encrp", "QM", "AQM", "DPA", "WFM", "WFMPro", "PM", "Speech", "Workflex", "Total"])
    ws.append(["Avaya LLC", "Acme Insurance Co.", "500-1999", "x", 900, "Perpetual", 0, 900, 0, 900, 0, 200, 900, 0, 0, 0, 216000, 216000])
    ws.append(["Cisco Systems, Inc.", "Northwind Health Plan", "2000-4999", "x", 2500, "SaaS", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0])
    ws.append(["Avaya LLC", "CIGNA Corporation", "5000-9999", "x", 5100, "SaaS", 0, 0, 0, 0, 0, 0, 5100, 0, 0, 0, 0, 0])
    ws.append(["ISI Telemanagement Solutions, LLC", "Bigco Utilities", "2000-4999", "x", 2100, "Term", 0, 0, 0, 0, 0, 300, 2100, 0, 0, 0, 0, 0])
    p2 = wb.create_sheet("Verint Partner Only 2024")
    p2.append(["Partner", "Account", "Max Verint", None, "Max Verint license", "Model", "DPA Seats", "Vertical Alignment", "DPA-Max Licence"])
    p2.append(["Avaya LLC", "Acme Insurance Co.", "500-1999", "x", 900, "Perpetual", 200, 1, 700])
    p2.append(["Cisco Systems, Inc.", "Northwind Health Plan", "2000-4999", "x", 2500, "SaaS", 0, 2, 2500])
    p2.append(["Direct", "CIGNA Corporation", None, None, None, "Verint Operations Manager - SaaS", 5100])
    p2.append(["Direct", "Ally Financial Inc.", None, None, None, "Verint Operations Visualizer (AAMT) - SaaS", 2306])
    p2.append(["TTEC", "VW Credit, Inc.", None, None, None, "Verint eg Work Manager", 350])
    p2.append(["Direct", "Five9, Inc", None, None, None, "Verint eg Work Manager", 6000])
    p = scratch / "selftest_whitespace.xlsx"
    wb.save(p)
    res = run(p, scratch / "selftest_out")
    print(json.dumps(res, indent=2))
    rows = list(csv.DictReader((scratch / "selftest_out" / "3xg_candidates_ranked.csv").open()))
    by = {r["account"]: r for r in rows}
    assert rows[0]["account"] == "Ally Financial Inc.", rows[0]
    assert by["Acme Insurance Co."]["bucket"] == "dpa_deployed" and by["Acme Insurance Co."]["wfm_in_footprint"] == "yes"
    assert by["Northwind Health Plan"]["bucket"] == "verint_no_dpa"
    assert by["VW Credit, Inc."]["bucket"] == "backoffice_deployment"
    assert by["CIGNA Corporation"]["flags"].startswith("HOLD"), by["CIGNA Corporation"]  # Cigna is a Salesforce customer
    assert by["Five9, Inc"]["flags"].startswith("HOLD") and "Partner" in by["Five9, Inc"]["flags"], by["Five9, Inc"]
    assert all(c not in rows[0] for c in ("Workflex", "Total", "QM", "Rec")), rows[0].keys()
    ext = list(csv.DictReader((scratch / "selftest_out" / "3xg_extended_pool.csv").open()))
    assert any(r["account"] == "Bigco Utilities" for r in ext), ext
    print("SELFTEST OK: OV first, buckets right, WFM footprint joined, Cigna held, dollar and module columns absent, extended pool built")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--file", type=pathlib.Path, help="workbook path (default: newest .xlsx in the inbox)")
    ap.add_argument("--inbox", type=pathlib.Path, default=DEFAULT_INBOX)
    ap.add_argument("--out", type=pathlib.Path, help="output folder (default: <inbox>/run_<date>)")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        selftest(pathlib.Path(os.environ.get("SCRATCH", "/tmp")) / "partner_intake_selftest")
        return
    path = a.file
    if not path:
        cands = sorted(a.inbox.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True) if a.inbox.exists() else []
        if not cands:
            sys.exit(f"no .xlsx in {a.inbox}. Save Frank's attachment there and rerun.")
        path = cands[0]
    res = run(path, a.out or (a.inbox / f"run_{dt.date.today().isoformat()}"))
    print(json.dumps(res, indent=2))


if __name__ == "__main__":
    main()
