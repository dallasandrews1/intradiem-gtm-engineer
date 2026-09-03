#!/usr/bin/env python3
"""Build the strategy .xlsx and .html map from a single strategy.json.

Usage:
    pip install openpyxl --break-system-packages   # if not installed
    python3 build_pack.py strategy.json output_dir/

Reads strategy.json (schema in references/strategy-json-spec.md) and writes:
    <output_dir>/<slug>_Meeting_Strategy.xlsx
    <output_dir>/<slug>_Strategy_Map.html

The HTML template must sit at ../assets/strategy_map_template.html relative
to this script (it does, inside the skill folder).
"""
import json
import html as h
import re
import sys
from pathlib import Path

BADGE_CLASSES = {"now": "b-now", "soon": "b-q3", "warm": "b-warm", "nurture": "b-nurture"}


def esc(s):
    return h.escape(str(s or ""))


def slugify(title):
    return re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")


# ---------------- XLSX ----------------

def build_xlsx(data, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    navy = "1F3864"
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor=navy)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(bottom=Side(style="thin", color="D9D9D9"))
    tier_fills = {"1": PatternFill("solid", fgColor="E8F7EF"),
                  "2": PatternFill("solid", fgColor="FDF6E3"),
                  "3": PatternFill("solid", fgColor="F3F4F6")}

    wb = Workbook()

    # Sheet 1: Strategy Tracker
    ws = wb.active
    ws.title = "Strategy Tracker"
    cols = ["Tier", "Account", "Est. agents", "M&A / org status", "Key signal (dated)",
            "Target personas", "Hook / opening angle", "Recommended play", "Timing", "Warmth / notes"]
    widths = [6, 22, 16, 30, 38, 30, 42, 42, 16, 28]
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = header_font, header_fill, Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for tier in data["tiers"]:
        t = str(tier["tier"])
        for a in tier["accounts"]:
            ws.append([t, a["name"], a.get("est_agents", ""), a.get("org_status", ""),
                       a.get("signal", ""), a.get("targets", ""), a.get("hook", ""),
                       a.get("play", ""), a.get("timing", ""), a.get("warmth", "")])
            r = ws.max_row
            for i in range(1, len(cols) + 1):
                cell = ws.cell(row=r, column=i)
                cell.alignment, cell.border = wrap, thin
            ws.cell(row=r, column=1).fill = tier_fills.get(t, tier_fills["3"])
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=2).font = Font(bold=True)

    # Sheet 2: Playbook
    pb = wb.create_sheet("Playbook")
    pb.column_dimensions["A"].width = 34
    pb.column_dimensions["B"].width = 90
    pb.append([data["title"].upper() + " - MEETING-BOOKING PLAYBOOK", ""])
    pb.cell(row=1, column=1).font = Font(bold=True, size=13, color=navy)
    pb.append(["", ""])
    reframe = data.get("reframe", {})
    if reframe:
        pb.append([reframe.get("heading", "THE BIG REFRAME").upper(), reframe.get("text", "")])
    for item in data.get("playbook", []):
        pb.append(["", ""])
        pb.append([item["title"].upper(), item["text"]])
    for row in pb.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # Sheet 3: Removed - Parked
    if data.get("parked"):
        pk = wb.create_sheet("Removed - Parked")
        pk.append(["Account", "Reason"])
        pk.column_dimensions["A"].width = 28
        pk.column_dimensions["B"].width = 60
        for c in pk[1]:
            c.font, c.fill = header_font, header_fill
        for p in data["parked"]:
            pk.append([p["account"], p["reason"]])
            pk.cell(row=pk.max_row, column=2).alignment = wrap

    wb.save(out_path)


# ---------------- HTML ----------------

def card_html(a, tier):
    badge_cls = BADGE_CLASSES.get(a.get("badge_type", "nurture"), "b-nurture")
    warm = a.get("warmnote")
    warm_html = f'<div class="lbl">Warmth</div><span class="warmnote">{esc(warm)}</span>' if warm else ""
    return f'''    <div class="card t{tier}"><div class="card-top" onclick="this.parentNode.classList.toggle('open')">
      <div class="row1"><span class="name">{esc(a["name"])}</span><span class="badge {badge_cls}">{esc(a.get("badge",""))}</span></div>
      <div class="meta">{esc(a.get("meta",""))}</div>
      <div class="signal"><b>Signal:</b> {esc(a.get("signal",""))}</div>
      <div class="chev">&#9662; full play</div></div>
      <div class="detail">
        <div class="lbl">Targets</div>{esc(a.get("targets",""))}
        <div class="lbl">Hook</div><div class="hook">{esc(a.get("hook",""))}</div>
        <div class="lbl">Play</div>{esc(a.get("play",""))}
        {warm_html}
      </div>
    </div>'''


def build_html(data, out_path, template_path):
    tpl = Path(template_path).read_text(encoding="utf-8")
    stats = "\n".join(
        f'  <div class="stat s{i+1}"><div class="n">{esc(s["n"])}</div><div class="l">{esc(s["label"])}</div></div>'
        for i, s in enumerate(data.get("stats", [])[:4]))
    tiers_html = []
    for tier in data["tiers"]:
        t = tier["tier"]
        cards = "\n".join(card_html(a, t) for a in tier["accounts"])
        tiers_html.append(f'''<div class="tier t{t}c">
  <div class="tier-head"><h2>{esc(tier["name"])}</h2><span class="tag">{esc(tier.get("tag",""))}</span></div>
  <div class="grid">
{cards}
  </div>
</div>''')
    playbook = "\n".join(
        f'    <div class="play"><h3><span>{p.get("icon","&#128204;")}</span>{esc(p["title"])}</h3>{esc(p["text"])}</div>'
        for p in data.get("playbook", []))
    reframe = data.get("reframe", {})
    out = (tpl.replace("{{TITLE}}", esc(data["title"]))
              .replace("{{SUBTITLE}}", esc(data.get("subtitle", "")))
              .replace("{{REFRAME_HEADING}}", esc(reframe.get("heading", "")))
              .replace("{{REFRAME_TEXT}}", esc(reframe.get("text", "")))
              .replace("{{STATS}}", stats)
              .replace("{{TIERS}}", "\n\n".join(tiers_html))
              .replace("{{PLAYBOOK}}", playbook)
              .replace("{{FOOTER}}", esc(data.get("footer", "Click any card for the full play. Internal use only."))))
    Path(out_path).write_text(out, encoding="utf-8")


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    data = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)
    slug = slugify(data["title"])
    template = Path(__file__).resolve().parent.parent / "assets" / "strategy_map_template.html"
    build_xlsx(data, out_dir / f"{slug}_Meeting_Strategy.xlsx")
    build_html(data, out_dir / f"{slug}_Strategy_Map.html", template)
    print(f"Wrote {slug}_Meeting_Strategy.xlsx and {slug}_Strategy_Map.html to {out_dir}")


if __name__ == "__main__":
    main()
